"""Módulo Historial de movimientos — solo admin."""
import io
from collections import Counter
from datetime import date, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse

from app.supabase_client import supabase
from app.core.historial import TABLA, ETIQUETAS
from app.core.fechas import ZONA_BOLIVIA, a_bolivia, limites_utc
from app.core.auth import solo, ADMIN

router = APIRouter(prefix="/historial", tags=["Historial"], dependencies=[Depends(solo(ADMIN))])

LIMITE_PANTALLA = 1000   # filas que se muestran en la tabla (el Excel trae todo)
PAGINA = 1000            # Supabase devuelve como máximo 1000 filas por consulta
MAXIMO = 20000           # tope de seguridad para el Excel

MODULOS = {"inventario": "Inventario", "compras": "Compras", "ventas": "Ventas"}


def _consultar(inicio: date, fin: date, usuario: Optional[str], accion: Optional[str]) -> list:
    desde_utc, hasta_utc = limites_utc(inicio, fin)
    filas, desde = [], 0
    while desde < MAXIMO:
        q = supabase.table(TABLA).select("*")\
            .gte("fecha", desde_utc)\
            .lte("fecha", hasta_utc)
        if usuario:
            q = q.eq("usuario", usuario)
        if accion:
            q = q.eq("accion", accion)
        lote = q.order("id", desc=True).range(desde, desde + PAGINA - 1).execute().data
        filas.extend(lote)
        if len(lote) < PAGINA:
            break
        desde += PAGINA
    return filas


def _resumen(filas: list) -> dict:
    por_accion = Counter(f["accion"] for f in filas)
    por_usuario = Counter(f["usuario"] for f in filas)
    roles = {f["usuario"]: f.get("rol") for f in filas}
    return {
        "por_accion": dict(por_accion),
        "por_usuario": [{"usuario": u, "rol": roles.get(u), "cantidad": c} for u, c in por_usuario.most_common()],
    }


def _hora_bolivia(iso: Optional[str]):
    """Hora de Bolivia sin zona horaria, que es lo que Excel entiende."""
    dt = a_bolivia(iso)
    return dt.replace(tzinfo=None) if dt else None


@router.get("/movimientos")
def listar_movimientos(inicio: date, fin: date, usuario: Optional[str] = None, accion: Optional[str] = None):
    filas = _consultar(inicio, fin, usuario, accion)
    return {
        "total": len(filas),
        "movimientos": filas[:LIMITE_PANTALLA],
        **_resumen(filas),
    }


@router.get("/usuarios")
def listar_usuarios():
    return supabase.table("usuarios").select("username, role").order("username").execute().data


@router.get("/excel")
def exportar_excel(inicio: date, fin: date, usuario: Optional[str] = None, accion: Optional[str] = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    filas = _consultar(inicio, fin, usuario, accion)
    if fin < inicio:
        inicio, fin = fin, inicio

    naranja = PatternFill(start_color="FF6B2B", end_color="FF6B2B", fill_type="solid")
    blanco_negrita = Font(bold=True, color="FFFFFF")
    arriba = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()

    # ── Hoja 1: detalle ──
    ws = wb.active
    ws.title = "Historial"
    ws.append(["Historial de movimientos — Almacen Cori"])
    ws["A1"].font = Font(bold=True, size=14, color="E85510")
    ahora = datetime.now(timezone.utc).astimezone(ZONA_BOLIVIA)
    ws.append([
        f"Período: {inicio:%d/%m/%Y} al {fin:%d/%m/%Y}  ·  "
        f"Usuario: {usuario or 'Todos'}  ·  Tipo: {ETIQUETAS.get(accion, 'Todos') if accion else 'Todos'}  ·  "
        f"Movimientos: {len(filas)}  ·  Generado: {ahora:%d/%m/%Y %H:%M}"
    ])
    ws["A2"].font = Font(italic=True, color="5C4A40")
    ws.append([])

    encabezados = ["Fecha y hora", "Usuario", "Rol", "Movimiento", "Módulo", "Producto", "Nivel", "Detalle"]
    ws.append(encabezados)
    fila_encabezado = ws.max_row
    for celda in ws[fila_encabezado]:
        celda.fill = naranja
        celda.font = blanco_negrita

    for f in filas:
        ws.append([
            _hora_bolivia(f.get("fecha")),
            f.get("usuario"),
            (f.get("rol") or "").capitalize(),
            ETIQUETAS.get(f.get("accion"), f.get("accion")),
            MODULOS.get(f.get("modulo"), f.get("modulo")),
            f.get("producto_nombre") or "",
            f.get("nivel") or "",
            f.get("detalle") or "",
        ])
        ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy hh:mm"
        for celda in ws[ws.max_row]:
            celda.alignment = arriba

    for letra, ancho in zip("ABCDEFGH", [17, 26, 10, 18, 12, 30, 12, 70]):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=1)
    if filas:
        ws.auto_filter.ref = f"A{fila_encabezado}:H{ws.max_row}"

    # ── Hoja 2: resumen por usuario ──
    ws2 = wb.create_sheet("Resumen por usuario")
    acciones = [a for a in ETIQUETAS if any(f["accion"] == a for f in filas)]
    ws2.append(["Usuario", "Rol"] + [ETIQUETAS[a] for a in acciones] + ["Total"])
    for celda in ws2[1]:
        celda.fill = naranja
        celda.font = blanco_negrita
    for u in _resumen(filas)["por_usuario"]:
        propias = [f for f in filas if f["usuario"] == u["usuario"]]
        conteo = Counter(f["accion"] for f in propias)
        ws2.append([u["usuario"], (u["rol"] or "").capitalize()] + [conteo.get(a, 0) for a in acciones] + [len(propias)])
    for col in ws2.columns:
        largo = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws2.column_dimensions[col[0].column_letter].width = min(largo + 3, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Historial_{inicio}_a_{fin}.xlsx"},
    )
