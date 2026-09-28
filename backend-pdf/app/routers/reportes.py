from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from app.supabase_client import supabase
from datetime import datetime, date
import io

router = APIRouter(prefix="/reportes", tags=["Reportes"])


def get_ventas_en_rango(inicio: str, fin: str):
    return supabase.table("ventas")\
        .select("*, clientes(nombre)")\
        .gte("fecha", f"{inicio}T00:00:00")\
        .lte("fecha", f"{fin}T23:59:59")\
        .eq("estado", "completada")\
        .execute()


def get_compras_en_rango(inicio: str, fin: str):
    return supabase.table("compras")\
        .select("*, proveedores(nombre)")\
        .gte("fecha", f"{inicio}T00:00:00")\
        .lte("fecha", f"{fin}T23:59:59")\
        .order("fecha", desc=True)\
        .execute()


def _monto_cobrado_ahora(v):
    total = v["total"]
    if v.get("es_credito"):
        return min(v.get("monto_recibido") or 0, total)
    return total


@router.get("/ventas/resumen")
def resumen_ventas(inicio: str, fin: str):
    result = get_ventas_en_rango(inicio, fin)
    ventas = result.data

    ingresos = sum(v["total"] for v in ventas)
    descuentos = sum(v.get("descuento") or 0 for v in ventas)
    cantidad = len(ventas)

    return {
        "total_transacciones": cantidad,
        "ingresos_totales": ingresos,
        "descuentos": descuentos,
        "ticket_promedio": ingresos / cantidad if cantidad > 0 else 0,
        "efectivo": sum(_monto_cobrado_ahora(v) for v in ventas if v.get("metodo_pago") == "efectivo"),
        "qr": sum(_monto_cobrado_ahora(v) for v in ventas if v.get("metodo_pago") in ["qr", "transferencia"]),
        "tarjeta": sum(_monto_cobrado_ahora(v) for v in ventas if v.get("metodo_pago") == "tarjeta"),
    }


@router.get("/ventas/lista")
def lista_ventas(inicio: str, fin: str):
    result = supabase.table("ventas")\
        .select("*, clientes(nombre)")\
        .gte("fecha", f"{inicio}T00:00:00")\
        .lte("fecha", f"{fin}T23:59:59")\
        .order("fecha", desc=True)\
        .execute()

    ventas = []
    for v in result.data:
        row = {**v}
        row["cliente_nombre"] = (v.get("clientes") or {}).get("nombre")
        row.pop("clientes", None)
        ventas.append(row)
    return ventas


@router.get("/ventas/top-productos")
def top_productos(inicio: str, fin: str):
    ventas = get_ventas_en_rango(inicio, fin)
    if not ventas.data:
        return []

    venta_ids = [v["id"] for v in ventas.data]
    detalles = supabase.table("detalle_ventas")\
        .select("*, productos(codigo, nombre)")\
        .in_("venta_id", venta_ids)\
        .execute()

    agrupado = {}
    for d in detalles.data:
        prod = d.get("productos") or {}
        pid = d["producto_id"]
        if pid not in agrupado:
            agrupado[pid] = {
                "codigo": prod.get("codigo", "—"),
                "nombre": prod.get("nombre", "Desconocido"),
                "unidades": 0,
                "ingresos": 0,
                "ganancia": 0
            }
        agrupado[pid]["unidades"] += d["cantidad"]
        agrupado[pid]["ingresos"] += d["subtotal"]
        costo = d.get("precio_compra", 0) * d["cantidad"]
        agrupado[pid]["ganancia"] += d["subtotal"] - costo

    resultado = sorted(agrupado.values(), key=lambda x: x["unidades"], reverse=True)
    return resultado[:20]


@router.get("/stock-critico")
def stock_critico():
    result = supabase.table("producto_niveles")\
        .select("nivel, stock, stock_minimo, productos(codigo, nombre, categorias(nombre))")\
        .eq("activo", True)\
        .execute()

    criticos = []
    for n in result.data:
        if n["stock_minimo"] > 0 and n["stock"] <= n["stock_minimo"]:
            prod = n.get("productos") or {}
            criticos.append({
                "codigo": prod.get("codigo", "—"),
                "nombre": prod.get("nombre", "Desconocido"),
                "nivel": n["nivel"],
                "categoria": (prod.get("categorias") or {}).get("nombre", "—"),
                "stock": n["stock"],
                "minimo": n["stock_minimo"],
                "faltante": n["stock_minimo"] - n["stock"]
            })
    return sorted(criticos, key=lambda x: x["stock"])


@router.get("/rentabilidad")
def rentabilidad(inicio: str, fin: str):
    ventas = get_ventas_en_rango(inicio, fin)
    if not ventas.data:
        return {"ganancia_neta": 0}

    venta_ids = [v["id"] for v in ventas.data]
    detalles = supabase.table("detalle_ventas")\
        .select("subtotal, precio_compra, cantidad")\
        .in_("venta_id", venta_ids)\
        .execute()

    ganancia = sum(
        d["subtotal"] - (d.get("precio_compra", 0) * d["cantidad"])
        for d in detalles.data
    )
    return {"ganancia_neta": ganancia}


@router.get("/compras/lista")
def lista_compras(inicio: str, fin: str):
    result = get_compras_en_rango(inicio, fin)
    compras = []
    for c in result.data:
        row = {**c}
        row["proveedor_nombre"] = (c.get("proveedores") or {}).get("nombre") or "Sin proveedor"
        row.pop("proveedores", None)
        compras.append(row)
    return compras


@router.get("/compras/resumen")
def resumen_compras(inicio: str, fin: str):
    result = get_compras_en_rango(inicio, fin)
    compras = result.data
    return {
        "total_compras": len(compras),
        "total_gastado": sum(c["total"] for c in compras)
    }


@router.get("/excel")
def exportar_excel(inicio: str, fin: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    ventas_result = get_ventas_en_rango(inicio, fin)
    ventas_data = ventas_result.data
    venta_ids = [v["id"] for v in ventas_data]

    detalles_ventas_data = []
    if venta_ids:
        detalles_ventas_data = supabase.table("detalle_ventas")\
            .select("*, ventas(numero_venta), productos(codigo, nombre), producto_niveles(nivel)")\
            .in_("venta_id", venta_ids).execute().data

    compras_result = get_compras_en_rango(inicio, fin)
    compras_data = compras_result.data
    compra_ids = [c["id"] for c in compras_data]

    detalles_compras_data = []
    if compra_ids:
        detalles_compras_data = supabase.table("detalle_compras")\
            .select("*, compras(numero_compra), productos(codigo, nombre), producto_niveles(nivel)")\
            .in_("compra_id", compra_ids).execute().data

    inventario_data = supabase.table("producto_niveles")\
        .select("*, productos(codigo, nombre, categorias(nombre))")\
        .eq("activo", True).execute().data

    top = top_productos(inicio, fin)

    abonos_data = supabase.table("cliente_movimientos")\
        .select("*, clientes(nombre)")\
        .eq("tipo", "abono")\
        .gte("fecha", f"{inicio}T00:00:00")\
        .lte("fecha", f"{fin}T23:59:59")\
        .order("fecha", desc=True)\
        .execute().data

    saldos_pendientes_data = supabase.table("clientes")\
        .select("*")\
        .gt("saldo_pendiente", 0)\
        .eq("activo", True)\
        .order("saldo_pendiente", desc=True)\
        .execute().data

    wb = Workbook()
    encabezado_fill = PatternFill(start_color="FF6B2B", end_color="FF6B2B", fill_type="solid")
    encabezado_font = Font(bold=True, color="FFFFFF")

    def escribir_encabezados(ws, encabezados):
        ws.append(encabezados)
        for cell in ws[1]:
            cell.fill = encabezado_fill
            cell.font = encabezado_font

    def autoajustar(ws):
        for col in ws.columns:
            largo = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(largo + 3, 40)

    # Hoja 1: Ventas
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Ventas"
    escribir_encabezados(ws1, ["N° Venta", "Fecha", "Cliente", "Crédito", "Método", "Subtotal", "Descuento", "Total", "Estado"])
    for v in ventas_data:
        cliente_nombre = (v.get("clientes") or {}).get("nombre") or "Consumidor final"
        metodo_mostrar = "Crédito" if v.get("es_credito") else v.get("metodo_pago")
        ws1.append([
            v["numero_venta"], v["fecha"][:10], cliente_nombre,
            "Sí" if v.get("es_credito") else "No",
            metodo_mostrar, v.get("subtotal"), v.get("descuento"), v["total"], v.get("estado")
        ])
    autoajustar(ws1)

    # Hoja 2: Detalle de ventas
    ws2 = wb.create_sheet("Detalle Ventas")
    escribir_encabezados(ws2, ["N° Venta", "Producto", "Nivel", "Cantidad", "Precio Unitario", "Subtotal"])
    for d in detalles_ventas_data:
        venta_info = d.get("ventas") or {}
        prod = d.get("productos") or {}
        nivel = d.get("producto_niveles") or {}
        ws2.append([
            venta_info.get("numero_venta"), prod.get("nombre"), nivel.get("nivel"),
            d["cantidad"], d["precio_unitario"], d["subtotal"]
        ])
    autoajustar(ws2)

    # Hoja 3: Top productos
    ws3 = wb.create_sheet("Top Productos")
    escribir_encabezados(ws3, ["Código", "Producto", "Unidades", "Ingresos", "Ganancia"])
    for p in top:
        ws3.append([p["codigo"], p["nombre"], p["unidades"], p["ingresos"], p["ganancia"]])
    autoajustar(ws3)

    # Hoja 4: Inventario actual (todos los niveles)
    ws4 = wb.create_sheet("Inventario")
    escribir_encabezados(ws4, ["Código", "Producto", "Categoría", "Nivel", "Precio Compra", "Precio Venta", "Stock", "Stock Mínimo"])
    for n in inventario_data:
        prod = n.get("productos") or {}
        cat = (prod.get("categorias") or {}).get("nombre", "—")
        ws4.append([
            prod.get("codigo"), prod.get("nombre"), cat, n["nivel"],
            n["precio_compra"], n["precio_venta"], n["stock"], n["stock_minimo"]
        ])
    autoajustar(ws4)

    # Hoja 5: Compras a proveedores
    ws5 = wb.create_sheet("Compras")
    escribir_encabezados(ws5, ["N° Compra", "Proveedor", "Fecha", "Total", "Notas"])
    for c in compras_data:
        proveedor_nombre = (c.get("proveedores") or {}).get("nombre") or "Sin proveedor"
        ws5.append([c["numero_compra"], proveedor_nombre, c["fecha"][:10], c["total"], c.get("notas", "")])
    autoajustar(ws5)

    # Hoja 6: Detalle de compras
    ws6 = wb.create_sheet("Detalle Compras")
    escribir_encabezados(ws6, ["N° Compra", "Producto", "Nivel", "Cantidad", "Precio Unitario", "Subtotal"])
    for d in detalles_compras_data:
        compra_info = d.get("compras") or {}
        prod = d.get("productos") or {}
        nivel = d.get("producto_niveles") or {}
        ws6.append([
            compra_info.get("numero_compra"), prod.get("nombre"), nivel.get("nivel"),
            d["cantidad"], d["precio_unitario"], d["subtotal"]
        ])
    autoajustar(ws6)

    # Hoja 7: Abonos de clientes (pagos a cuenta durante el período)
    ws7 = wb.create_sheet("Abonos de Clientes")
    escribir_encabezados(ws7, ["Cliente", "Fecha", "Monto", "Motivo"])
    for a in abonos_data:
        cliente_nombre = (a.get("clientes") or {}).get("nombre") or "—"
        ws7.append([cliente_nombre, a["fecha"][:10], a["monto"], a.get("motivo", "")])
    autoajustar(ws7)

    # Hoja 8: Saldos pendientes actuales (foto del momento, no filtrado por fecha)
    ws8 = wb.create_sheet("Cuentas por Cobrar")
    escribir_encabezados(ws8, ["Cliente", "Teléfono", "Límite de Crédito", "Saldo Pendiente"])
    for c in saldos_pendientes_data:
        ws8.append([c["nombre"], c.get("telefono", ""), c["limite_credito"], c["saldo_pendiente"]])
    autoajustar(ws8)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Reporte_{inicio}_a_{fin}.xlsx"}
    )